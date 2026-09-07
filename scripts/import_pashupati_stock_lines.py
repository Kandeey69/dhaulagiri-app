from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import shutil
import sqlite3
import sys
import uuid
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl


APP_ID = "com.easysolution.businesssuite"
COMPANY_ID = "pashupati-chemicals-2082-83"
COMPANY_NAME = "Pashupati Chemicals"
FISCAL_YEAR_ID = f"{COMPANY_ID}-2082-83"
WORKBOOK_PATH = Path(r"C:\Users\Kandeey\Documents\Pashupati stock.xlsx")


def stock_database_safe_company_id(company_id: str) -> str:
    normalized = company_id.strip()
    if not normalized or normalized == "default":
        return ""

    slug = "".join(character.lower() if character.isalnum() else "-" for character in normalized)
    while "--" in slug:
        slug = slug.replace("--", "-")
    slug = slug.strip("-")[:32] or "company"

    hash_value = 0xCBF29CE484222325
    prime = 0x100000001B3
    mask = 0xFFFFFFFFFFFFFFFF
    for character in normalized:
        hash_value ^= ord(character)
        hash_value = (hash_value * prime) & mask
    return f"{slug}-{hash_value:016x}"


def appdata_root() -> Path:
    return Path(os.environ["APPDATA"]) / APP_ID


def stock_db_path() -> Path:
    return appdata_root() / f"inventorytracked-stock-{stock_database_safe_company_id(COMPANY_ID)}.db"


def base_stock_db_path() -> Path:
    return appdata_root() / "inventorytracked-stock.db"


def accounts_db_path() -> Path:
    return appdata_root() / f"accounts-{COMPANY_ID}.db"


def purchase_db_path() -> Path:
    return appdata_root() / f"import-purchases-{COMPANY_ID}.db"


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime | dt.date):
        return f"{value.year:04d}/{value.month:02d}/{value.day:02d}"
    text = normalize_text(value).replace("-", "/")
    parts = text.split("/")
    if len(parts) == 3 and all(part.isdigit() for part in parts):
        return f"{int(parts[0]):04d}/{int(parts[1]):02d}/{int(parts[2]):02d}"
    return text


def number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, int | float):
        return float(value)
    return float(str(value).replace(",", "").strip())


def money(value: float) -> float:
    return round(float(value), 2)


def item_code(name: str) -> str:
    cleaned = []
    previous_dash = False
    for character in name.upper():
        if character.isalnum():
            cleaned.append(character)
            previous_dash = False
        elif not previous_dash:
            cleaned.append("-")
            previous_dash = True
    return "".join(cleaned).strip("-")[:48] or "STOCK-ITEM"


@dataclass
class SalesLine:
    bill_no: str
    date_bs: str
    party_name: str
    item_name: str
    quantity: float
    rate: float
    total: float


@dataclass
class PurchaseLine:
    bill_no: str
    date_bs: str
    party_name: str
    item_name: str
    quantity: float
    entry_rate: float
    entry_amount: float
    landed_cost: float


def read_sheet_records(workbook: openpyxl.Workbook, sheet_name: str, header_row: int) -> list[dict[str, Any]]:
    ws = workbook[sheet_name]
    headers = [normalize_text(cell) for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    records: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        record = {headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    return records


def read_workbook() -> tuple[list[SalesLine], list[PurchaseLine], list[str]]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    notes: list[str] = []
    if "stock type" in workbook.sheetnames:
        ws = workbook["stock type"]
        for row in ws.iter_rows(values_only=True):
            for value in row:
                text = normalize_text(value)
                if text and text.lower() not in {
                    "china clay powder (pyrophyllite)",
                    "talcum powder (201)",
                    "talcum powder (st 850 koohinoor)",
                    "talcum powder (grade s)",
                    "talcum powder (381)",
                    "talcum powder (kohinoor)",
                    "talcum powder (700)",
                }:
                    notes.append(text)

    sales_lines = []
    for record in read_sheet_records(workbook, "stock register", 1):
        sales_lines.append(
            SalesLine(
                bill_no=normalize_text(record.get("Bill Number")),
                date_bs=normalize_date(record.get("Date")),
                party_name=normalize_text(record.get("Party Name")),
                item_name=normalize_text(record.get("Particulars")),
                quantity=number(record.get("Quantity")),
                rate=number(record.get("Rate")),
                total=money(number(record.get("Total"))),
            )
        )

    purchase_lines = []
    for record in read_sheet_records(workbook, "Purchase register", 2):
        if not normalize_text(record.get("Bill Number")):
            continue
        purchase_lines.append(
            PurchaseLine(
                bill_no=normalize_text(record.get("Bill Number")),
                date_bs=normalize_date(record.get("Date")),
                party_name=normalize_text(record.get("Party Name")),
                item_name=normalize_text(record.get("Particulars")),
                quantity=number(record.get("Quantity")),
                entry_rate=number(record.get("INR Rate")),
                entry_amount=money(number(record.get("Line total INR"))),
                landed_cost=money(number(record.get("Landed cost NPR"))),
            )
        )

    return sales_lines, purchase_lines, notes


def fetch_rows(path: Path, query: str) -> list[dict[str, Any]]:
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    try:
        return [dict(row) for row in con.execute(query)]
    finally:
        con.close()


def load_sales_bills() -> dict[str, dict[str, Any]]:
    rows = fetch_rows(
        accounts_db_path(),
        """
        SELECT sales.*, parties.name AS party_name
        FROM sales
        LEFT JOIN parties ON parties.id = sales.party_id
        """,
    )
    return {normalize_text(row["bill_no"]): row for row in rows}


def load_purchase_bills() -> dict[str, dict[str, Any]]:
    rows = fetch_rows(
        purchase_db_path(),
        """
        SELECT import_purchases.*, parties.name AS vendor_name
        FROM import_purchases
        LEFT JOIN parties ON parties.id = import_purchases.vendorPartyId
        """,
    )
    return {normalize_text(row["vendorBillNumber"]): row for row in rows}


def group_by_bill(lines: list[SalesLine] | list[PurchaseLine]) -> dict[str, list[Any]]:
    grouped: dict[str, list[Any]] = defaultdict(list)
    for line in lines:
        grouped[line.bill_no].append(line)
    return dict(grouped)


def compare_name(left: str, right: str) -> bool:
    return " ".join(left.lower().split()) == " ".join(right.lower().split())


def verify() -> tuple[dict[str, Any], list[SalesLine], list[PurchaseLine], dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    sales_lines, purchase_lines, notes = read_workbook()
    sales_bills = load_sales_bills()
    purchase_bills = load_purchase_bills()
    sales_by_bill = group_by_bill(sales_lines)
    purchase_by_bill = group_by_bill(purchase_lines)

    errors: list[str] = []
    warnings: list[str] = []

    for bill_no, bill in sales_bills.items():
        if bill_no not in sales_by_bill:
            errors.append(f"Sales bill {bill_no} exists in app but has no stock-register line.")
            continue
        lines = sales_by_bill[bill_no]
        line_total = money(sum(line.total for line in lines))
        bill_total = money(number(bill["sales_amount"]))
        if abs(line_total - bill_total) > 0.01:
            errors.append(f"Sales bill {bill_no} line total {line_total:.2f} does not match sales amount {bill_total:.2f}.")
        if lines[0].date_bs != normalize_date(bill["date_bs"]):
            errors.append(f"Sales bill {bill_no} date {lines[0].date_bs} does not match app date {normalize_date(bill['date_bs'])}.")
        if not compare_name(lines[0].party_name, normalize_text(bill["party_name"])):
            errors.append(f"Sales bill {bill_no} party {lines[0].party_name!r} does not match app party {normalize_text(bill['party_name'])!r}.")

    for bill_no in sales_by_bill:
        if bill_no not in sales_bills:
            errors.append(f"Stock-register sales bill {bill_no} was not found in app sales.")

    for bill_no, bill in purchase_bills.items():
        if bill_no not in purchase_by_bill:
            errors.append(f"Purchase bill {bill_no} exists in app but has no purchase-register line.")
            continue
        lines = purchase_by_bill[bill_no]
        line_total = money(sum(line.entry_amount for line in lines))
        bill_total = money(number(bill["amountIC"]))
        landed_total = money(sum(line.landed_cost for line in lines))
        bill_landed = money(number(bill["landedCostNPR"]))
        if abs(line_total - bill_total) > 0.01:
            errors.append(f"Purchase bill {bill_no} INR line total {line_total:.2f} does not match amountIC {bill_total:.2f}.")
        if abs(landed_total - bill_landed) > 0.05:
            errors.append(f"Purchase bill {bill_no} landed-cost line total {landed_total:.2f} does not match app landed cost {bill_landed:.2f}.")
        source_date = normalize_date(bill.get("debitNoteDate") or bill.get("billDate"))
        if lines[0].date_bs != source_date:
            warnings.append(f"Purchase bill {bill_no} register date {lines[0].date_bs} differs from app source date {source_date}.")
        if not compare_name(lines[0].party_name, normalize_text(bill["vendor_name"])):
            errors.append(f"Purchase bill {bill_no} party {lines[0].party_name!r} does not match app vendor {normalize_text(bill['vendor_name'])!r}.")

    for bill_no in purchase_by_bill:
        if bill_no not in purchase_bills:
            errors.append(f"Purchase-register bill {bill_no} was not found in app purchases.")

    sales_note_bills = ["118", "174", "208", "209"]
    note_checks = {
        bill_no: {
            "in_sales_app": bill_no in sales_bills,
            "in_sales_register": bill_no in sales_by_bill,
            "sales_line_total": money(sum(line.total for line in sales_by_bill.get(bill_no, []))),
            "sales_amount": money(number(sales_bills[bill_no]["sales_amount"])) if bill_no in sales_bills else None,
            "items": [line.item_name for line in sales_by_bill.get(bill_no, [])],
        }
        for bill_no in sales_note_bills
    }

    report = {
        "workbook": str(WORKBOOK_PATH),
        "accounts_db": str(accounts_db_path()),
        "purchase_db": str(purchase_db_path()),
        "stock_db": str(stock_db_path()),
        "sales_line_count": len(sales_lines),
        "sales_bill_count_in_register": len(sales_by_bill),
        "sales_bill_count_in_app": len(sales_bills),
        "purchase_line_count": len(purchase_lines),
        "purchase_bill_count_in_register": len(purchase_by_bill),
        "purchase_bill_count_in_app": len(purchase_bills),
        "stock_items": sorted({line.item_name for line in [*sales_lines, *purchase_lines]}),
        "notes_from_workbook": notes,
        "note_bill_checks": note_checks,
        "warnings": warnings,
        "errors": errors,
    }
    return report, sales_lines, purchase_lines, sales_bills, purchase_bills


def init_stock_db(con: sqlite3.Connection) -> None:
    con.executescript(
        """
        PRAGMA busy_timeout = 10000;
        PRAGMA journal_mode = WAL;
        PRAGMA synchronous = NORMAL;
        CREATE TABLE IF NOT EXISTS stock_items (
          id TEXT PRIMARY KEY,
          code TEXT NOT NULL UNIQUE,
          name TEXT NOT NULL,
          unit TEXT NOT NULL,
          opening_qty REAL DEFAULT 0,
          opening_rate REAL DEFAULT 0,
          reorder_level REAL DEFAULT 0,
          is_active INTEGER DEFAULT 1,
          created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS stock_purchase_bills (
          id TEXT PRIMARY KEY,
          bill_no TEXT NOT NULL,
          date_bs TEXT NOT NULL,
          supplier_name TEXT NOT NULL,
          source TEXT NOT NULL,
          source_type TEXT NOT NULL DEFAULT 'Local Purchase',
          reference_no TEXT,
          remarks TEXT,
          created_at TEXT NOT NULL,
          source_amount REAL,
          source_amount_npr REAL,
          source_currency TEXT,
          source_exchange_rate REAL,
          source_fiscal_year_id TEXT,
          source_grand_total REAL,
          source_landed_cost_npr REAL,
          source_lifecycle_status TEXT
        );
        CREATE TABLE IF NOT EXISTS stock_purchase_lines (
          id TEXT PRIMARY KEY,
          bill_id TEXT NOT NULL,
          item_id TEXT NOT NULL,
          quantity REAL NOT NULL,
          rate REAL NOT NULL,
          amount REAL NOT NULL,
          entry_rate REAL,
          entry_amount REAL,
          FOREIGN KEY (bill_id) REFERENCES stock_purchase_bills(id),
          FOREIGN KEY (item_id) REFERENCES stock_items(id)
        );
        CREATE TABLE IF NOT EXISTS stock_sales_bills (
          id TEXT PRIMARY KEY,
          bill_no TEXT NOT NULL,
          date_bs TEXT NOT NULL,
          customer_name TEXT NOT NULL,
          source_type TEXT NOT NULL DEFAULT 'Sale',
          remarks TEXT,
          created_at TEXT NOT NULL,
          source_amount REAL,
          source_amount_npr REAL,
          source_currency TEXT,
          source_exchange_rate REAL,
          source_fiscal_year_id TEXT,
          source_grand_total REAL,
          source_landed_cost_npr REAL,
          source_lifecycle_status TEXT
        );
        CREATE TABLE IF NOT EXISTS stock_sales_lines (
          id TEXT PRIMARY KEY,
          bill_id TEXT NOT NULL,
          item_id TEXT NOT NULL,
          quantity REAL NOT NULL,
          rate REAL NOT NULL,
          amount REAL NOT NULL,
          FOREIGN KEY (bill_id) REFERENCES stock_sales_bills(id),
          FOREIGN KEY (item_id) REFERENCES stock_items(id)
        );
        CREATE UNIQUE INDEX IF NOT EXISTS idx_stock_items_code ON stock_items(code);
        CREATE INDEX IF NOT EXISTS idx_stock_purchase_bills_bill_no ON stock_purchase_bills(bill_no);
        CREATE INDEX IF NOT EXISTS idx_stock_purchase_bills_source_type ON stock_purchase_bills(source_type, id);
        CREATE INDEX IF NOT EXISTS idx_stock_sales_bills_bill_no ON stock_sales_bills(bill_no);
        CREATE INDEX IF NOT EXISTS idx_stock_sales_bills_source_type ON stock_sales_bills(source_type, id);
        """
    )


def choose_sales_lines_for_import(
    bill_no: str,
    lines: list[SalesLine],
    source_amount: float,
    corrections: list[str],
    unresolved: list[str],
) -> list[SalesLine]:
    target = money(source_amount)
    current = money(sum(line.total for line in lines))
    if abs(current - target) <= 0.01:
        return lines

    if len(lines) <= 10:
        best_subset: list[SalesLine] | None = None
        for mask in range(1, 1 << len(lines)):
            subset = [line for index, line in enumerate(lines) if mask & (1 << index)]
            subset_total = money(sum(line.total for line in subset))
            if abs(subset_total - target) <= 0.01:
                if best_subset is None or len(subset) > len(best_subset):
                    best_subset = subset
        if best_subset is not None and len(best_subset) < len(lines):
            corrections.append(
                f"Sales bill {bill_no}: ignored {len(lines) - len(best_subset)} duplicate line(s) because a subset matches the source amount."
            )
            return best_subset

    difference = money(target - current)
    if abs(difference) <= 1:
        adjusted = list(lines)
        last = adjusted[-1]
        adjusted[-1] = SalesLine(
            bill_no=last.bill_no,
            date_bs=last.date_bs,
            party_name=last.party_name,
            item_name=last.item_name,
            quantity=last.quantity,
            rate=round((last.total + difference) / last.quantity, 6) if last.quantity else last.rate,
            total=money(last.total + difference),
        )
        corrections.append(f"Sales bill {bill_no}: adjusted last line by {difference:.2f} for rounding.")
        return adjusted

    unresolved.append(
        f"Sales bill {bill_no}: entered workbook line total {current:.2f}, but source amount is {target:.2f}."
    )
    return lines


def choose_purchase_lines_for_import(
    bill_no: str,
    lines: list[PurchaseLine],
    source_amount: float,
    source_landed_cost: float,
    corrections: list[str],
    unresolved: list[str],
) -> list[PurchaseLine]:
    target = money(source_amount)
    current = money(sum(line.entry_amount for line in lines))
    landed_total = money(sum(line.landed_cost for line in lines))
    if abs(current - target) <= 0.01:
        return lines

    if len(lines) == 1 and abs(landed_total - money(source_landed_cost)) <= 0.05:
        line = lines[0]
        corrected = PurchaseLine(
            bill_no=line.bill_no,
            date_bs=line.date_bs,
            party_name=line.party_name,
            item_name=line.item_name,
            quantity=line.quantity,
            entry_rate=round(target / line.quantity, 6) if line.quantity else line.entry_rate,
            entry_amount=target,
            landed_cost=line.landed_cost,
        )
        corrections.append(
            f"Purchase bill {bill_no}: corrected entry INR total from {current:.2f} to source amount {target:.2f}; landed cost was already matched."
        )
        return [corrected]

    unresolved.append(
        f"Purchase bill {bill_no}: entered workbook INR total {current:.2f}, but source amount is {target:.2f}."
    )
    return lines


def apply_import(strict: bool = True) -> dict[str, Any]:
    report, sales_lines, purchase_lines, sales_bills, purchase_bills = verify()
    if strict and report["errors"]:
        raise SystemExit(json.dumps(report, indent=2, ensure_ascii=False))

    now = dt.datetime.now(dt.UTC).isoformat().replace("+00:00", "Z")
    corrections: list[str] = []
    unresolved: list[str] = []
    raw_sales_by_bill = group_by_bill(sales_lines)
    raw_purchase_by_bill = group_by_bill(purchase_lines)
    sales_by_bill: dict[str, list[SalesLine]] = {}
    purchase_by_bill: dict[str, list[PurchaseLine]] = {}

    for bill_no, lines in raw_sales_by_bill.items():
        if bill_no not in sales_bills:
            unresolved.append(f"Sales bill {bill_no}: workbook line exists but source bill was not found.")
            continue
        bill = sales_bills[bill_no]
        sales_by_bill[bill_no] = choose_sales_lines_for_import(
            bill_no,
            lines,
            number(bill["sales_amount"]),
            corrections,
            unresolved,
        )

    for bill_no, bill in sales_bills.items():
        if bill_no not in raw_sales_by_bill:
            unresolved.append(f"Sales bill {bill_no}: source bill has no stock-register line.")

    for bill_no, lines in raw_purchase_by_bill.items():
        if bill_no not in purchase_bills:
            unresolved.append(f"Purchase bill {bill_no}: workbook line exists but source bill was not found.")
            continue
        bill = purchase_bills[bill_no]
        purchase_by_bill[bill_no] = choose_purchase_lines_for_import(
            bill_no,
            lines,
            number(bill["amountIC"]),
            number(bill["landedCostNPR"]),
            corrections,
            unresolved,
        )

    for bill_no, bill in purchase_bills.items():
        if bill_no not in raw_purchase_by_bill:
            unresolved.append(f"Purchase bill {bill_no}: source bill has no purchase-register line.")

    items = sorted({line.item_name for lines in [*sales_by_bill.values(), *purchase_by_bill.values()] for line in lines})
    item_ids = {name: str(uuid.uuid5(uuid.NAMESPACE_URL, f"{COMPANY_ID}:stock-item:{name}")) for name in items}

    path = stock_db_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    backup_path = ""
    if path.exists():
        backup = path.with_suffix(path.suffix + f".bak-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}")
        shutil.copy2(path, backup)
        backup_path = str(backup)
    con = sqlite3.connect(path)
    try:
        init_stock_db(con)
        con.execute("BEGIN IMMEDIATE")
        con.execute("DELETE FROM stock_purchase_lines")
        con.execute("DELETE FROM stock_sales_lines")
        con.execute("DELETE FROM stock_purchase_bills")
        con.execute("DELETE FROM stock_sales_bills")
        con.execute("DELETE FROM stock_items")

        used_codes: set[str] = set()
        for name in items:
            code = item_code(name)
            base_code = code
            suffix = 2
            while code in used_codes:
                code = f"{base_code[:44]}-{suffix}"
                suffix += 1
            used_codes.add(code)
            con.execute(
                """
                INSERT INTO stock_items (
                  id, code, name, unit, opening_qty, opening_rate, reorder_level, is_active, created_at
                ) VALUES (?, ?, ?, 'KG', 0, 0, 0, 1, ?)
                """,
                (item_ids[name], code, name, now),
            )

        for bill_no, lines in purchase_by_bill.items():
            bill = purchase_bills[bill_no]
            bill_id = f"Import Purchase:{bill['id']}"
            con.execute(
                """
                INSERT INTO stock_purchase_bills (
                  id, bill_no, date_bs, supplier_name, source, source_type, reference_no, remarks, created_at,
                  source_amount, source_amount_npr, source_currency, source_exchange_rate, source_fiscal_year_id,
                  source_grand_total, source_landed_cost_npr, source_lifecycle_status
                ) VALUES (?, ?, ?, ?, 'Importation', 'Import Purchase', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    bill_id,
                    bill_no,
                    normalize_date(bill.get("debitNoteDate") or bill.get("billDate")) or lines[0].date_bs,
                    lines[0].party_name,
                    normalize_text(bill.get("debitNoteNumber")),
                    "Imported from Pashupati stock.xlsx",
                    now,
                    number(bill.get("amountIC")),
                    number(bill.get("supplierAmountNPR")),
                    normalize_text(bill.get("supplierCurrency")) or "INR",
                    number(bill.get("supplierExchangeRate")),
                    normalize_text(bill.get("fiscalYearId")) or FISCAL_YEAR_ID,
                    number(bill.get("landedCostNPR")),
                    number(bill.get("landedCostNPR")),
                    normalize_text(bill.get("lifecycleStatus")) or "POSTED",
                ),
            )
            for index, line in enumerate(lines, start=1):
                rate = round(line.landed_cost / line.quantity, 6) if line.quantity else 0
                con.execute(
                    """
                    INSERT INTO stock_purchase_lines (
                      id, bill_id, item_id, quantity, rate, amount, entry_rate, entry_amount
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(uuid.uuid5(uuid.NAMESPACE_URL, f"{bill_id}:line:{index}:{line.item_name}")),
                        bill_id,
                        item_ids[line.item_name],
                        line.quantity,
                        rate,
                        line.landed_cost,
                        line.entry_rate,
                        line.entry_amount,
                    ),
                )

        for bill_no, lines in sales_by_bill.items():
            bill = sales_bills[bill_no]
            bill_id = bill["id"]
            con.execute(
                """
                INSERT INTO stock_sales_bills (
                  id, bill_no, date_bs, customer_name, source_type, remarks, created_at,
                  source_amount, source_amount_npr, source_currency, source_exchange_rate, source_fiscal_year_id,
                  source_grand_total, source_landed_cost_npr, source_lifecycle_status
                ) VALUES (?, ?, ?, ?, 'Sale', ?, ?, ?, ?, 'NPR', NULL, ?, ?, NULL, ?)
                """,
                (
                    bill_id,
                    bill_no,
                    lines[0].date_bs,
                    lines[0].party_name,
                    "Imported from Pashupati stock.xlsx",
                    now,
                    number(bill.get("sales_amount")),
                    number(bill.get("sales_amount")),
                    normalize_text(bill.get("fiscal_year_id")) or FISCAL_YEAR_ID,
                    number(bill.get("total_amount")),
                    normalize_text(bill.get("lifecycle_status")) or "POSTED",
                ),
            )
            for index, line in enumerate(lines, start=1):
                con.execute(
                    """
                    INSERT INTO stock_sales_lines (
                      id, bill_id, item_id, quantity, rate, amount
                    ) VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(uuid.uuid5(uuid.NAMESPACE_URL, f"{bill_id}:line:{index}:{line.item_name}")),
                        bill_id,
                        item_ids[line.item_name],
                        line.quantity,
                        line.rate,
                        line.total,
                    ),
                )

        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()

    if not base_stock_db_path().exists():
        sqlite3.connect(base_stock_db_path()).close()

    verify_report = verify_stock_db(path)
    report["applied"] = True
    report["strict"] = strict
    report["backup_path"] = backup_path
    report["corrections_applied"] = corrections
    report["unresolved_after_import"] = unresolved
    report["stock_db_verification"] = verify_report
    return report


def verify_stock_db(path: Path) -> dict[str, Any]:
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    try:
        counts = {
            table: con.execute(f"SELECT COUNT(*) AS count FROM {table}").fetchone()["count"]
            for table in [
                "stock_items",
                "stock_purchase_bills",
                "stock_purchase_lines",
                "stock_sales_bills",
                "stock_sales_lines",
            ]
        }
        sales_total = con.execute("SELECT ROUND(COALESCE(SUM(amount), 0), 2) AS total FROM stock_sales_lines").fetchone()["total"]
        purchase_total = con.execute("SELECT ROUND(COALESCE(SUM(amount), 0), 2) AS total FROM stock_purchase_lines").fetchone()["total"]
        entry_total = con.execute("SELECT ROUND(COALESCE(SUM(entry_amount), 0), 2) AS total FROM stock_purchase_lines").fetchone()["total"]
        return {
            "counts": counts,
            "sales_line_total": sales_total,
            "purchase_landed_cost_total": purchase_total,
            "purchase_entry_inr_total": entry_total,
        }
    finally:
        con.close()


def workbook_detail_rows(sheet_name: str, header_row: int, bill_header: str, bill_numbers: list[str]) -> list[dict[str, Any]]:
    values_workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    formula_workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False, read_only=True)
    ws = values_workbook[sheet_name]
    formula_ws = formula_workbook[sheet_name]
    headers = [normalize_text(cell) for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    bill_index = headers.index(bill_header) + 1
    wanted = set(bill_numbers)
    rows: list[dict[str, Any]] = []

    for row_number in range(header_row + 1, ws.max_row + 1):
        bill_no = normalize_text(ws.cell(row_number, bill_index).value)
        if bill_no not in wanted:
            continue
        values = {}
        formulas = {}
        for column_number, header in enumerate(headers, start=1):
            if not header:
                continue
            value = ws.cell(row_number, column_number).value
            formula_value = formula_ws.cell(row_number, column_number).value
            values[header] = normalize_date(value) if "date" in header.lower() or header == "Date" else value
            if isinstance(formula_value, str) and formula_value.startswith("="):
                formulas[header] = formula_value
        rows.append({"row": row_number, "values": values, "formulas": formulas})

    return rows


def detail_report() -> dict[str, Any]:
    sales_bills = ["53", "68", "174", "209", "254"]
    purchase_bills = ["2026/27/73"]
    con = sqlite3.connect(accounts_db_path())
    con.row_factory = sqlite3.Row
    try:
        sales_db = [
            dict(row)
            for row in con.execute(
                """
                SELECT sales.id, bill_no, date_bs, parties.name AS party_name,
                       sales_amount, vat_amount, total_amount
                FROM sales
                LEFT JOIN parties ON parties.id = sales.party_id
                WHERE bill_no IN ('53', '68', '174', '209', '254')
                ORDER BY CAST(bill_no AS INTEGER)
                """
            )
        ]
    finally:
        con.close()

    con = sqlite3.connect(purchase_db_path())
    con.row_factory = sqlite3.Row
    try:
        purchase_db = [
            dict(row)
            for row in con.execute(
                """
                SELECT import_purchases.id, vendorBillNumber, billDate,
                       parties.name AS vendor_name, amountIC, supplierAmountNPR,
                       agentServiceBillDate, landedCostNPR
                FROM import_purchases
                LEFT JOIN parties ON parties.id = import_purchases.vendorPartyId
                WHERE vendorBillNumber = '2026/27/73'
                """
            )
        ]
    finally:
        con.close()

    return {
        "workbook": {
            "stock register": workbook_detail_rows("stock register", 1, "Bill Number", sales_bills),
            "Pashupati sales": workbook_detail_rows("Pashupati sales", 1, "billNo", sales_bills),
            "Purchase register": workbook_detail_rows("Purchase register", 2, "Bill Number", purchase_bills),
            "Pashupati purchases": workbook_detail_rows("Pashupati purchases", 1, "vendorBillNumber", purchase_bills),
        },
        "sales_db": sales_db,
        "purchase_db": purchase_db,
    }


def final_reconciliation() -> dict[str, Any]:
    stock_path = stock_db_path()
    con = sqlite3.connect(stock_path)
    con.row_factory = sqlite3.Row
    try:
        stock_counts = {
            table: con.execute(f"SELECT COUNT(*) AS count FROM {table}").fetchone()["count"]
            for table in [
                "stock_items",
                "stock_purchase_bills",
                "stock_purchase_lines",
                "stock_sales_bills",
                "stock_sales_lines",
            ]
        }
        stock = {
            "path": str(stock_path),
            "counts": stock_counts,
            "items": [dict(row) for row in con.execute("SELECT code, name, unit FROM stock_items ORDER BY name")],
            "sales_total": con.execute("SELECT ROUND(COALESCE(SUM(amount), 0), 2) AS total FROM stock_sales_lines").fetchone()["total"],
            "purchase_entry_inr_total": con.execute("SELECT ROUND(COALESCE(SUM(entry_amount), 0), 2) AS total FROM stock_purchase_lines").fetchone()["total"],
            "purchase_landed_total": con.execute("SELECT ROUND(COALESCE(SUM(amount), 0), 2) AS total FROM stock_purchase_lines").fetchone()["total"],
            "bill_174": [
                dict(row)
                for row in con.execute(
                    """
                    SELECT b.bill_no, b.date_bs, b.customer_name, i.name AS item_name,
                           l.quantity, l.rate, l.amount
                    FROM stock_sales_bills b
                    JOIN stock_sales_lines l ON l.bill_id = b.id
                    JOIN stock_items i ON i.id = l.item_id
                    WHERE b.bill_no = '174'
                    """
                )
            ],
            "bill_53_count": con.execute("SELECT COUNT(*) AS count FROM stock_sales_bills WHERE bill_no = '53'").fetchone()["count"],
            "bill_209_count": con.execute("SELECT COUNT(*) AS count FROM stock_sales_bills WHERE bill_no = '209'").fetchone()["count"],
        }
    finally:
        con.close()

    con = sqlite3.connect(accounts_db_path())
    con.row_factory = sqlite3.Row
    try:
        sales = {
            "count": con.execute("SELECT COUNT(*) AS count FROM sales").fetchone()["count"],
            "sales_total": con.execute("SELECT ROUND(COALESCE(SUM(sales_amount), 0), 2) AS total FROM sales").fetchone()["total"],
        }
    finally:
        con.close()

    con = sqlite3.connect(purchase_db_path())
    con.row_factory = sqlite3.Row
    try:
        purchases = {
            "count": con.execute("SELECT COUNT(*) AS count FROM import_purchases").fetchone()["count"],
            "entry_inr_total": con.execute("SELECT ROUND(COALESCE(SUM(amountIC), 0), 2) AS total FROM import_purchases").fetchone()["total"],
            "landed_total": con.execute("SELECT ROUND(COALESCE(SUM(landedCostNPR), 0), 2) AS total FROM import_purchases").fetchone()["total"],
        }
    finally:
        con.close()

    return {"stock": stock, "source_sales": sales, "source_purchases": purchases}


def purchase_status_reconciliation() -> dict[str, Any]:
    def values_match(left: Any, right: Any, tolerance: float = 0.5) -> bool:
        return abs(number(left) - number(right)) <= tolerance

    purchase_con = sqlite3.connect(purchase_db_path())
    purchase_con.row_factory = sqlite3.Row
    stock_con = sqlite3.connect(stock_db_path())
    stock_con.row_factory = sqlite3.Row

    try:
        purchases = [
            dict(row)
            for row in purchase_con.execute(
                """
                SELECT import_purchases.*, parties.name AS vendor_name
                FROM import_purchases
                LEFT JOIN parties ON parties.id = import_purchases.vendorPartyId
                """
            )
        ]
        statuses = []
        for purchase in purchases:
            stock_bill_id = f"Import Purchase:{purchase['id']}"
            stock_bill = stock_con.execute(
                "SELECT * FROM stock_purchase_bills WHERE id = ? AND source_type = 'Import Purchase'",
                (stock_bill_id,),
            ).fetchone()
            if not stock_bill:
                statuses.append({
                    "bill": purchase["vendorBillNumber"],
                    "status": "Pending",
                    "reason": "Missing stock purchase bill.",
                })
                continue

            lines = [
                dict(row)
                for row in stock_con.execute(
                    "SELECT * FROM stock_purchase_lines WHERE bill_id = ?",
                    (stock_bill_id,),
                )
            ]
            entry_total = money(sum(number(line.get("entry_amount") if line.get("entry_amount") is not None else line.get("amount")) for line in lines))
            landed_total = money(sum(number(line.get("amount")) for line in lines))
            problems = []
            source_date = normalize_date(purchase.get("debitNoteDate") or purchase.get("billDate"))

            if normalize_text(stock_bill["bill_no"]) != normalize_text(purchase["vendorBillNumber"]):
                problems.append("bill number")
            if normalize_date(stock_bill["date_bs"]) != source_date:
                problems.append("date")
            if not compare_name(stock_bill["supplier_name"], purchase.get("vendor_name") or ""):
                problems.append("vendor")
            if not values_match(entry_total, purchase["amountIC"]):
                problems.append(f"entry total {entry_total:.2f} vs {money(number(purchase['amountIC'])):.2f}")
            if not values_match(landed_total, purchase["landedCostNPR"]):
                problems.append(f"landed total {landed_total:.2f} vs {money(number(purchase['landedCostNPR'])):.2f}")
            if not values_match(stock_bill["source_amount"], purchase["amountIC"]):
                problems.append("source amount")
            if not values_match(stock_bill["source_amount_npr"], purchase["supplierAmountNPR"]):
                problems.append("source amount NPR")
            if normalize_text(stock_bill["source_currency"]) != normalize_text(purchase["supplierCurrency"]):
                problems.append("currency")
            if not values_match(stock_bill["source_exchange_rate"], purchase["supplierExchangeRate"], 0.000001):
                problems.append("exchange rate")
            if normalize_text(stock_bill["source_fiscal_year_id"]) != normalize_text(purchase["fiscalYearId"]):
                problems.append("fiscal year")
            if not values_match(stock_bill["source_grand_total"], purchase["landedCostNPR"]):
                problems.append("grand total snapshot")
            if not values_match(stock_bill["source_landed_cost_npr"], purchase["landedCostNPR"]):
                problems.append("landed snapshot")
            if normalize_text(stock_bill["source_lifecycle_status"]) != normalize_text(purchase["lifecycleStatus"]):
                problems.append("lifecycle")

            statuses.append({
                "bill": purchase["vendorBillNumber"],
                "status": "Entered" if not problems else "Mismatch",
                "line_count": len(lines),
                "entry_total": entry_total,
                "landed_total": landed_total,
                "problems": problems,
            })
    finally:
        purchase_con.close()
        stock_con.close()

    return {
        "purchase_bills": len(statuses),
        "entered": sum(1 for status in statuses if status["status"] == "Entered"),
        "mismatch": sum(1 for status in statuses if status["status"] == "Mismatch"),
        "pending": sum(1 for status in statuses if status["status"] == "Pending"),
        "bad": [status for status in statuses if status["status"] != "Entered"],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--apply-verified", action="store_true")
    parser.add_argument("--details", action="store_true")
    parser.add_argument("--summary", action="store_true")
    parser.add_argument("--purchase-status", action="store_true")
    args = parser.parse_args()

    if args.purchase_status:
        report = purchase_status_reconciliation()
    elif args.summary:
        report = final_reconciliation()
    elif args.details:
        report = detail_report()
    elif args.apply:
        report = apply_import(strict=True)
    elif args.apply_verified:
        report = apply_import(strict=False)
    else:
        report = verify()[0]

    print(json.dumps(report, indent=2, ensure_ascii=False, default=str))
    return 0 if not report.get("errors") else 1


if __name__ == "__main__":
    sys.exit(main())
